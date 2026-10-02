"""
생성형 AI(Gemini)와 Python 기반 화학실험 데이터 정제 및 이상치 분석 자동화 파이프라인

실행:  python src/analyzer.py
입력:  data/experiment_raw.csv
출력:  output/yield_analysis_plot.png, output/analysis_result.xlsx
"""

import json
import os
import re
import sys
import time
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # GUI 없는 환경에서도 그래프 저장 가능
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from dotenv import load_dotenv

# ---------------------------------------------------------------------------
# 경로 및 상수 설정
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent.parent
DATA_PATH = BASE_DIR / "data" / "experiment_raw.csv"
OUTPUT_DIR = BASE_DIR / "output"
PLOT_PATH = OUTPUT_DIR / "yield_analysis_plot.png"
EXCEL_PATH = OUTPUT_DIR / "analysis_result.xlsx"

NUMERIC_COLS = ["Temp_C", "Pressure_bar", "Catalyst_g", "Reaction_Time_min", "Yield_pct", "Purity_pct"]
FEATURE_COLS = ["Temp_C", "Pressure_bar", "Catalyst_g", "Reaction_Time_min"]

# 공정 운전 허용 범위 (물리적 타당성 검사용)
OPERATING_RANGE = {
    "Temp_C": (140, 220),
    "Pressure_bar": (1.0, 3.0),
    "Catalyst_g": (0.2, 1.5),
    "Reaction_Time_min": (30, 120),
    "Yield_pct": (0, 100),
    "Purity_pct": (0, 100),
}

# 분류·요약 작업은 경량(Lite) 모델로 충분하고, Lite 모델이 서버 과부하(503)가 훨씬 적음
DEFAULT_MODEL = "gemini-3.5-flash-lite"  # 버전 고정 → 실행할 때마다 같은 모델로 재현성 확보
FALLBACK_MODELS = ["gemini-flash-lite-latest", "gemini-3.1-flash-lite"]  # 기본 모델 실패 시 순서대로 시도

CATEGORIES = ["정상", "촉매 문제", "온도 이상", "압력 누출", "기타"]

# 환각 검증용 키워드 사전 (AI 분류와 교차 검증)
KEYWORD_RULES = {
    "촉매 문제": ["촉매"],
    "온도 이상": ["온도", "오버슈트", "냉각", "열폭주", "과열"],
    "압력 누출": ["압력 누출", "누출", "가스켓", "밸브"],
}
ISSUE_WORDS = ["누락", "손실", "혼입", "불량", "오류", "부반응"]


# ---------------------------------------------------------------------------
# 1. 환경변수 로드
# ---------------------------------------------------------------------------
def load_api_key():
    """.env 에서 GEMINI_API_KEY 를 읽는다. 없으면 안내 후 None 반환 (오프라인 모드)."""
    load_dotenv(BASE_DIR / ".env")
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key or api_key == "your_gemini_api_key_here":
        print("[안내] GEMINI_API_KEY 가 설정되지 않았습니다.")
        print("       .env.example 을 복사해 .env 를 만들고 API 키를 입력하세요.")
        print("       → 지금은 규칙 기반(키워드) 분류로 대체하여 파이프라인을 계속 실행합니다.\n")
        return None
    return api_key


# ---------------------------------------------------------------------------
# 2. 데이터 로드 및 전처리
# ---------------------------------------------------------------------------
REQUIRED_COLS = ["Experiment_ID", *NUMERIC_COLS, "Notes"]
PLOT_PALETTE = {"정상": "#2a6fdb", "이상치": "#d64545", "trend": "#2a6fdb"}


def load_and_clean(path):
    if not path.exists():
        raise FileNotFoundError(f"원본 데이터가 없습니다: {path}")
    raw = pd.read_csv(path, encoding="utf-8")
    print(f"[1] 원본 데이터 로드: {len(raw)}행 x {raw.shape[1]}열")
    return clean_data(raw)


def clean_data(raw):
    """원본 DataFrame 을 받아 결측치 처리·이상치 판정을 수행 (CLI·웹 공용)."""
    missing_cols = [c for c in REQUIRED_COLS if c not in raw.columns]
    if missing_cols:
        raise ValueError(f"필수 컬럼이 없습니다: {', '.join(missing_cols)}")
    if raw.empty:
        raise ValueError("데이터가 비어 있습니다.")

    raw = raw[REQUIRED_COLS].copy()
    df = raw.copy()

    # 결측치 현황 기록
    missing_report = df.isna().sum().rename("Missing_Count").to_frame()
    missing_report["Missing_Pct"] = (missing_report["Missing_Count"] / len(df) * 100).round(1)
    print("    결측치 현황:\n" + missing_report[missing_report["Missing_Count"] > 0].to_string())

    df[NUMERIC_COLS] = df[NUMERIC_COLS].apply(pd.to_numeric, errors="coerce")

    # 결측 위치를 먼저 플래그로 남긴 뒤 보정 (추적성 확보)
    flags = []
    for _, row in df.iterrows():
        missing = [c for c in NUMERIC_COLS if pd.isna(row[c])]
        flags.append("결측보정: " + ", ".join(missing) if missing else "")
    df["Data_Flag"] = flags

    # 독립변수(운전조건)는 중앙값 대체 - 이상치에 강건
    for col in FEATURE_COLS:
        df[col] = df[col].fillna(df[col].median())

    # 종속변수(수율)는 임의 대체하지 않고 분석 제외 표시 (데이터 왜곡 방지)
    df["Use_For_Analysis"] = df["Yield_pct"].notna()

    df["Notes"] = df["Notes"].fillna("기록 없음").astype(str).str.strip()

    # 이상치 탐지 ① 운전 범위 이탈
    range_issue = []
    for _, row in df.iterrows():
        out = [c for c, (lo, hi) in OPERATING_RANGE.items() if pd.notna(row[c]) and not lo <= row[c] <= hi]
        range_issue.append(", ".join(out))
    df["Range_Violation"] = range_issue

    # 이상치 탐지 ② 수율 IQR (Q1 - 1.5IQR 미만)
    valid_yield = df.loc[df["Use_For_Analysis"], "Yield_pct"]
    if len(valid_yield) >= 4:
        q1, q3 = valid_yield.quantile([0.25, 0.75])
        lower = q1 - 1.5 * (q3 - q1)
    else:  # 표본이 너무 적으면 IQR 판정 생략
        lower = -np.inf
    df["Yield_IQR_Outlier"] = df["Yield_pct"] < lower

    # 이상치 탐지 ③ 연구자가 Notes 에 직접 표기한 [이상치]
    df["Noted_Outlier"] = df["Notes"].str.contains(r"\[이상치\]", regex=True)

    df["Is_Outlier"] = df["Yield_IQR_Outlier"] | df["Noted_Outlier"] | (df["Range_Violation"] != "")

    print(f"    수율 IQR 하한: {lower:.1f}%  |  이상치 판정: {int(df['Is_Outlier'].sum())}건")
    return raw, df, missing_report


def summarize(df):
    normal = df[df["Use_For_Analysis"] & ~df["Is_Outlier"]]
    summary = df.loc[df["Use_For_Analysis"], NUMERIC_COLS].describe().T.round(2)
    summary["mean_excl_outlier"] = normal[NUMERIC_COLS].mean().round(2)

    corr = normal[NUMERIC_COLS].corr()["Yield_pct"].drop("Yield_pct").round(3)
    summary["corr_with_yield(정상군)"] = corr
    print("[2] 기초 통계 요약 완료")
    print(summary[["mean", "std", "min", "max"]].to_string())
    return summary


# ---------------------------------------------------------------------------
# 3. 시각화
# ---------------------------------------------------------------------------
def set_korean_font():
    candidates = ["AppleGothic", "Malgun Gothic", "NanumGothic", "Noto Sans CJK KR"]
    available = {f.name for f in matplotlib.font_manager.fontManager.ttflist}
    for name in candidates:
        if name in available:
            plt.rcParams["font.family"] = name
            break
    plt.rcParams["axes.unicode_minus"] = False


def make_plot(df, path, palette=None):
    """path 는 파일 경로 또는 BytesIO. palette 로 점·추세선 색을 바꿀 수 있음 (웹용)."""
    palette = palette or PLOT_PALETTE
    set_korean_font()
    sns.set_theme(style="whitegrid", font=plt.rcParams["font.family"])
    plt.rcParams["axes.unicode_minus"] = False

    plot_df = df[df["Use_For_Analysis"]].copy()
    plot_df["구분"] = np.where(plot_df["Is_Outlier"], "이상치", "정상")
    normal = plot_df[~plot_df["Is_Outlier"]]

    point_colors = {k: palette[k] for k in ("정상", "이상치")}
    can_fit = len(normal) >= 3

    fig, axes = plt.subplots(1, 2, figsize=(13, 5.2))
    for ax, x, label in [
        (axes[0], "Temp_C", "반응온도 (°C)"),
        (axes[1], "Catalyst_g", "촉매 투입량 (g)"),
    ]:
        sns.scatterplot(data=plot_df, x=x, y="Yield_pct", hue="구분", palette=point_colors,
                        style="구분", markers={"정상": "o", "이상치": "X"},  # 색 + 모양으로 이중 구분
                        hue_order=["정상", "이상치"], style_order=["정상", "이상치"],
                        s=85, edgecolor="white", ax=ax)
        # 추세선은 정상 데이터만으로 적합 (이상치가 경향을 왜곡하지 않도록)
        if can_fit:
            sns.regplot(data=normal, x=x, y="Yield_pct", scatter=False, order=2 if x == "Temp_C" else 1,
                        color=palette["trend"], line_kws={"lw": 1.5, "ls": "--"}, ax=ax)
        for _, r in plot_df[plot_df["Is_Outlier"]].iterrows():
            ax.annotate(r["Experiment_ID"], (r[x], r["Yield_pct"]), xytext=(5, -10),
                        textcoords="offset points", fontsize=8, color=palette["이상치"])
        ax.set_xlabel(label)
        ax.set_ylabel("수율 Yield (%)")
        ax.set_title(f"{label.split(' (')[0]} vs 수율")
        ax.legend(title="")

    fig.suptitle("반응 조건별 수율 분석 (점선: 정상군 추세)", fontsize=14)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    if isinstance(path, Path):
        print(f"[3] 그래프 저장: {path.relative_to(BASE_DIR)}")


# ---------------------------------------------------------------------------
# 4. Gemini API 연동 - 비정형 Notes 분류 및 원인 요약
# ---------------------------------------------------------------------------
PROMPT_TEMPLATE = """너는 화학공학 반응기 공정 전문가다.
아래는 반응기 실험의 특이사항(Notes) 기록과 해당 실험의 정량 데이터다.

[작업]
각 실험마다 다음을 수행하라.
1. category: 반드시 다음 5개 중 하나만 선택 → {categories}
   - 정상: 공정상 문제 없음(기록 누락·계측 기록 오류만 있고 반응 자체는 정상인 경우 포함)
   - 촉매 문제: 촉매 불활성화, 촉매량 부족/계량 오류 등
   - 온도 이상: 온도 오버슈트, 냉각 실패, 열폭주 등
   - 압력 누출: 누출, 가스켓/밸브 불량으로 압력 유지 실패
   - 기타: 원료 오염, 샘플 손실, 계측기 이상 등 위에 해당하지 않는 경우
2. cause_summary: 수율(또는 순도) 저하 원인을 한국어 1문장으로 요약. 정상이면 "특이사항 없음".

[규칙]
- Notes 와 정량 데이터에 적힌 사실만 근거로 판단하고, 기록에 없는 원인을 추측해 만들지 마라.
- 근거가 불충분하면 category 는 "기타", cause_summary 는 "기록상 원인 불명확"으로 답하라.
- 출력은 설명 없이 JSON 배열만 반환하라.
  형식: [{{"Experiment_ID": "EXP_001", "category": "정상", "cause_summary": "특이사항 없음"}}, ...]

[데이터]
{records}
"""


def build_records(df):
    cols = ["Experiment_ID", "Temp_C", "Pressure_bar", "Catalyst_g", "Yield_pct", "Purity_pct", "Notes"]
    records = df[cols].replace({np.nan: None}).to_dict(orient="records")
    return json.dumps(records, ensure_ascii=False, indent=1)


def parse_json_array(text):
    """모델 응답에서 JSON 배열만 추출 (```json 코드블록 등 대응)."""
    match = re.search(r"\[.*\]", text, re.DOTALL)
    if not match:
        raise ValueError("응답에서 JSON 배열을 찾지 못했습니다.")
    return json.loads(match.group(0))


def classify_with_gemini(df, api_key):
    # 회사망 보안 프록시 등으로 SSL 인증서 오류가 나는 경우를 막기 위해
    # Python 기본 인증서 대신 OS(맥 키체인/Windows 인증서 저장소)의 인증서를 사용
    try:
        import truststore

        truststore.inject_into_ssl()
    except ImportError:
        pass

    from google import genai
    from google.genai import errors as genai_errors
    from google.genai import types

    model_name = os.getenv("GEMINI_MODEL", DEFAULT_MODEL)
    client = genai.Client(api_key=api_key, http_options=types.HttpOptions(timeout=120_000))  # ms
    prompt = PROMPT_TEMPLATE.format(categories=" / ".join(CATEGORIES), records=build_records(df))

    config = types.GenerateContentConfig(
        temperature=0.1,
        response_mime_type="application/json",
        automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),  # 함수 호출 미사용
    )

    # 503(서버 과부하)·429(요청 한도)는 잠시 후 재시도, 404(모델 종료)는 바로 다음 모델로 전환
    models = list(dict.fromkeys([model_name, *FALLBACK_MODELS]))
    max_attempts = 2
    last_error = None
    print(f"[4] Gemini API 호출 중... ({len(df)}건 일괄 처리)")
    for model in models:
        for attempt in range(1, max_attempts + 1):
            try:
                response = client.models.generate_content(model=model, contents=prompt, config=config)
                items = parse_json_array(response.text)
                print(f"    응답 성공 (model={model})")
                return {item["Experiment_ID"]: item for item in items}, model
            except genai_errors.APIError as e:
                last_error = e
                if e.code == 404:
                    print(f"    {model}: 사용할 수 없는 모델 → 다음 모델로 전환")
                    break
                if e.code not in (429, 500, 503):
                    raise
                if attempt < max_attempts:
                    print(f"    {model}: 서버 혼잡({e.code}) → 2초 후 재시도")
                    time.sleep(2)
                else:
                    print(f"    {model}: 서버 혼잡({e.code}) → 다음 모델로 전환")
    raise last_error


def rule_based_category(note):
    """키워드 기반 분류 - API 미사용 시 대체 및 AI 결과 교차검증 기준."""
    if "[이상치]" in note:
        # 압력 → 온도 → 촉매 순으로 판정
        for category in ["압력 누출", "온도 이상", "촉매 문제"]:
            if any(k in note for k in KEYWORD_RULES[category]):
                return category
        return "기타"
    # [이상치] 표기가 없으면 계측/기록 이슈 여부만 확인 ("온도 기록 누락. 반응은 정상" → 정상)
    if "반응은 정상" in note or not any(k in note for k in ISSUE_WORDS):
        return "정상"
    return "기타"


def rule_based_summary(row, category):
    if category == "정상":
        return "특이사항 없음"
    first = re.split(r"[.]", row["Notes"].replace("[이상치]", "").strip())[0].strip()
    return f"주요 원인(원본 기록): {first}."


def apply_ai_analysis(df, api_key):
    df = df.copy()
    df["Rule_Category"] = df["Notes"].apply(rule_based_category)

    ai_results, source = {}, "Rule-based (API 미사용)"
    if api_key:
        try:
            ai_results, used_model = classify_with_gemini(df, api_key)
            source = f"Gemini ({used_model})"
        except Exception as e:  # 네트워크/인증/파싱 오류 시에도 파이프라인은 끝까지 수행
            print(f"[경고] Gemini API 호출 실패 → 규칙 기반 분류로 대체합니다. ({type(e).__name__}: {e})")

    categories, summaries = [], []
    for _, row in df.iterrows():
        item = ai_results.get(row["Experiment_ID"])
        if item and item.get("category") in CATEGORIES:
            categories.append(item["category"])
            summaries.append(str(item.get("cause_summary", "")).strip())
        else:
            cat = row["Rule_Category"]
            categories.append(cat)
            summaries.append(rule_based_summary(row, cat))

    df["AI_Category"] = categories
    df["AI_Cause_Summary"] = summaries
    df["AI_Source"] = source

    # --- 환각(Hallucination) 검증 ---------------------------------------------
    # 기본값 True 로 두고, 원본 데이터와 모순되는 경우만 False 로 전환
    df["AI_Verified"] = True
    reasons = [[] for _ in range(len(df))]
    for i, (_, row) in enumerate(df.iterrows()):
        # ① 규칙 기반 분류와 불일치
        if row["AI_Category"] != row["Rule_Category"]:
            reasons[i].append(f"규칙분류({row['Rule_Category']})와 불일치")
        # ② 원본 Notes 에 [이상치] 표기가 있는데 AI 가 '정상'으로 판정
        if row["Noted_Outlier"] and row["AI_Category"] == "정상":
            reasons[i].append("원본 [이상치] 표기와 모순")
        # ③ 요약문에 원본에 없는 수치가 등장 (수치 환각)
        source_text = row["Notes"] + " " + " ".join(str(row[c]) for c in NUMERIC_COLS)
        for num in re.findall(r"\d+(?:\.\d+)?", row["AI_Cause_Summary"]):
            if num not in source_text:
                reasons[i].append(f"원본에 없는 수치 '{num}' 포함")
                break
    df["AI_Verified"] = [not r for r in reasons]
    df["Verify_Note"] = ["; ".join(r) if r else "원본 대조 일치" for r in reasons]

    n_fail = int((~df["AI_Verified"]).sum())
    print(f"    AI 분류 완료 (source: {source}) | 검증 불일치: {n_fail}건 → 엑셀에서 수동 검토 필요")
    return df


# ---------------------------------------------------------------------------
# 5. 엑셀 리포트 저장
# ---------------------------------------------------------------------------
def category_summary(df):
    return (
        df[df["Use_For_Analysis"]]
        .groupby("AI_Category")
        .agg(건수=("Experiment_ID", "count"), 평균수율=("Yield_pct", "mean"), 평균순도=("Purity_pct", "mean"))
        .reindex(CATEGORIES).dropna(how="all").round(2).reset_index()
    )


def save_excel(df, raw, summary, missing_report, path):
    from openpyxl.styles import Alignment, Font, PatternFill

    result_cols = [
        "Experiment_ID", *NUMERIC_COLS, "Notes", "AI_Category", "AI_Cause_Summary", "AI_Verified",
        "Verify_Note", "Rule_Category", "Is_Outlier", "Range_Violation", "Data_Flag", "Use_For_Analysis", "AI_Source",
    ]
    result = df[result_cols]

    cat_summary = category_summary(df)

    red = PatternFill("solid", fgColor="FDE2E2")
    yellow = PatternFill("solid", fgColor="FFF4CC")
    header_fill = PatternFill("solid", fgColor="1F4E78")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="Cleaned_Data", index=False)
        summary.to_excel(writer, sheet_name="Summary_Stats")
        cat_summary.to_excel(writer, sheet_name="AI_Category_Summary", index=False)
        missing_report.to_excel(writer, sheet_name="Missing_Report")
        raw.to_excel(writer, sheet_name="Raw_Data", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "B2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width * 1.3), 60)

        ws = writer.sheets["Cleaned_Data"]
        col_idx = {name: i + 1 for i, name in enumerate(result_cols)}
        for r, (_, row) in enumerate(result.iterrows(), start=2):
            if row["Is_Outlier"]:
                for c in range(1, len(result_cols) + 1):
                    ws.cell(row=r, column=c).fill = red
            if not row["AI_Verified"]:
                ws.cell(row=r, column=col_idx["AI_Verified"]).fill = yellow
                ws.cell(row=r, column=col_idx["Verify_Note"]).fill = yellow

    if isinstance(path, Path):
        print(f"[5] 엑셀 리포트 저장: {path.relative_to(BASE_DIR)}")


# ---------------------------------------------------------------------------
def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    api_key = load_api_key()

    raw, df, missing_report = load_and_clean(DATA_PATH)
    summary = summarize(df)
    make_plot(df, PLOT_PATH)
    df = apply_ai_analysis(df, api_key)

    try:
        save_excel(df, raw, summary, missing_report, EXCEL_PATH)
    except PermissionError:
        print(f"[오류] {EXCEL_PATH.name} 파일이 열려 있습니다. 엑셀을 닫고 다시 실행하세요.")
        sys.exit(1)

    print("\n완료! output/ 폴더에서 결과를 확인하세요.")


if __name__ == "__main__":
    main()
